"""
PRISM Performance Analysis - Word Document (.docx) Generator
Generates a structured, executive-grade DOCX report containing all performance
parameters, test execution matrices, endpoint breakdown, resource stats, and SLO evaluations.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_header_row(row, hex_color="1A365D"):
    for cell in row.cells:
        set_cell_background(cell, hex_color)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

def style_data_row(row, is_alternate=False, align_left_col0=True):
    bg_color = "F7FAFC" if is_alternate else "FFFFFF"
    for idx, cell in enumerate(row.cells):
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        for p in cell.paragraphs:
            if idx == 0 and align_left_col0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif idx == 1 and align_left_col0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(45, 55, 72)

def build_docx_report(output_path):
    doc = Document()
    
    # Page Setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = title.add_run("PRISM BACKEND LOAD TESTING")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(26, 54, 93)

    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run("Performance Analysis Parameters & Evaluation Matrix")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(15)
    r_sub.font.color.rgb = RGBColor(74, 85, 104)

    meta = doc.add_paragraph()
    meta_run = meta.add_run("Target: http://127.0.0.1:8080 (Spring Boot 4.x / Java 21) | Date: October 27, 2026 | Engine: Locust 2.47")
    meta_run.font.size = Pt(9.5)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(113, 128, 150)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 1. Executive Summary
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Executive Summary")
    r_h1.font.color.rgb = RGBColor(26, 54, 93)

    p_exec = doc.add_paragraph(
        "This document provides the formal performance analysis parameters and verification matrix for the PRISM "
        "Web3 Security Platform backend. Testing was executed in automated stages ranging from a 5-user smoke validation "
        "up to a high-scale 100-user saturation test executing 50,000 live requests. The PRISM Spring Boot container "
        "demonstrated outstanding stability, processing 50.71 req/s at an average latency of 4.86 ms (p95: 9.0 ms, p99: 11.0 ms) "
        "with zero 5xx server errors, zero database pool exhaustion, and negligible JVM heap growth."
    )
    p_exec.style.font.size = Pt(10)

    # -------------------------------------------------------------
    # 2. Test Execution Parameters Matrix
    # -------------------------------------------------------------
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Test Execution Parameters Matrix")
    r_h2.font.color.rgb = RGBColor(26, 54, 93)

    t1_headers = ["Scenario", "Purpose", "Users", "Spawn Rate", "Duration", "Shape", "Auth Mix", "Tripwires"]
    t1_data = [
        ["Scenario A: Smoke", "Pipeline & sanity check", "5", "2 u/s", "45 s", "Flat", "20% Pub / 80% Auth", ">5% err, >2000ms"],
        ["Scenario B: Baseline", "Reference performance", "10", "2 u/s", "3 min (180s)", "Flat", "20% Pub / 80% Auth", ">5% err, >2000ms"],
        ["Scenario C: Stepped Load", "Stage capacity check", "10->25->50", "2 u/s", "5 min (300s)", "Stepped (3 stg)", "20% Pub / 80% Auth", ">5% err, >2000ms"],
        ["100-User Pilot", "Ramp-up preflight check", "100", "10 u/s", "45 s", "Fast Ramp", "20% Pub / 80% Auth", ">5% err, >2000ms"],
        ["50k High-Scale Run", "Sustained capacity & saturation", "100", "10 u/s", "986 s (~16.4m)", "Sustained Peak", "20% Pub / 80% Auth", ">5% err, >2000ms"],
    ]

    t1 = doc.add_table(rows=len(t1_data)+1, cols=len(t1_headers))
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(t1_headers):
        t1.rows[0].cells[idx].text = name
    style_header_row(t1.rows[0])

    for row_idx, row_values in enumerate(t1_data):
        row = t1.rows[row_idx + 1]
        for col_idx, val in enumerate(row_values):
            row.cells[col_idx].text = val
        style_data_row(row, is_alternate=(row_idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 3. Cross-Scenario Performance Comparison Matrix
    # -------------------------------------------------------------
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Cross-Scenario Performance Comparison Matrix")
    r_h3.font.color.rgb = RGBColor(26, 54, 93)

    t2_headers = ["Metric", "Smoke (5 Users)", "Baseline (10 Users)", "Stepped (50 Users)", "Pilot (100 Users)", "50k Scale Run (100 Users)"]
    t2_data = [
        ["Total Requests", "111", "897", "4,863", "2,056", "50,000"],
        ["Validated Successes", "111 (100%)", "897 (100%)", "4,863 (100%)", "2,056 (100%)", "49,872 (99.74%)"],
        ["Total Failures", "0 (0.00%)", "0 (0.00%)", "0 (0.00%)", "0 (0.00%)", "128 (0.256%) *404 race"],
        ["Server 5xx Errors", "0 (0.00%)", "0 (0.00%)", "0 (0.00%)", "0 (0.00%)", "0 (0.00%)"],
        ["Throughput (RPS)", "2.52 req/s", "5.02 req/s", "16.20 req/s", "46.44 req/s", "50.71 req/s"],
        ["Latency Min", "3.14 ms", "2.33 ms", "1.66 ms", "1.43 ms", "1.28 ms"],
        ["Latency Avg", "8.71 ms", "7.73 ms", "5.78 ms", "4.86 ms", "4.86 ms"],
        ["Latency Median (p50)", "8.00 ms", "7.00 ms", "5.00 ms", "4.00 ms", "4.00 ms"],
        ["Latency p90", "14.00 ms", "12.00 ms", "9.00 ms", "8.00 ms", "8.00 ms"],
        ["Latency p95", "17.00 ms", "14.00 ms", "10.00 ms", "9.00 ms", "9.00 ms"],
        ["Latency p99", "30.00 ms", "18.00 ms", "13.00 ms", "13.00 ms", "11.00 ms"],
        ["Latency Max", "34.60 ms", "31.58 ms", "36.18 ms", "22.97 ms", "68.28 ms"],
        ["SLO Compliance", "PASS", "PASS", "PASS", "PASS", "PASS"],
    ]

    t2 = doc.add_table(rows=len(t2_data)+1, cols=len(t2_headers))
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(t2_headers):
        t2.rows[0].cells[idx].text = name
    style_header_row(t2.rows[0])

    for row_idx, row_values in enumerate(t2_data):
        row = t2.rows[row_idx + 1]
        for col_idx, val in enumerate(row_values):
            row.cells[col_idx].text = val
        style_data_row(row, is_alternate=(row_idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 4. 50,000-Request Benchmark Endpoint Breakdown Matrix
    # -------------------------------------------------------------
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. 50,000-Request Benchmark Endpoint Breakdown Matrix")
    r_h4.font.color.rgb = RGBColor(26, 54, 93)

    t3_headers = ["Method", "Endpoint / Action", "Requests", "Failures", "RPS", "Avg (ms)", "p50 (ms)", "p95 (ms)", "p99 (ms)", "Max (ms)", "Status"]
    t3_data = [
        ["GET", "/api/overview", "5,952", "0", "6.04", "6.41", "6.0", "10.0", "12.0", "32.72", "PASS"],
        ["GET", "/api/contracts", "5,117", "0", "5.19", "6.61", "6.0", "11.0", "13.0", "39.15", "PASS"],
        ["GET", "/actuator/health [Public]", "4,092", "0", "4.15", "3.92", "3.0", "7.0", "10.0", "29.12", "PASS"],
        ["GET", "/api/system/status", "4,087", "0", "4.14", "3.76", "3.0", "7.0", "8.0", "23.82", "PASS"],
        ["GET", "/api/findings", "3,981", "0", "4.04", "4.80", "4.0", "8.0", "9.0", "32.44", "PASS"],
        ["GET", "/api/contracts/[id]", "3,058", "32", "3.10", "4.76", "4.0", "8.0", "10.0", "32.89", "404 Race"],
        ["GET", "/api/runtime/events", "2,973", "0", "3.02", "3.95", "3.0", "7.0", "9.0", "68.28", "PASS"],
        ["GET", "/api/alerts", "2,962", "0", "3.00", "5.41", "5.0", "9.0", "10.0", "24.24", "PASS"],
        ["GET", "/api/runtime/telemetry", "2,955", "37", "3.00", "4.25", "4.0", "7.0", "9.0", "26.42", "404 Race"],
        ["GET", "/api/contracts/[id]/risk", "2,952", "30", "2.99", "4.85", "4.0", "8.0", "9.0", "52.80", "404 Race"],
        ["GET", "/api/overview [Expected 401]", "2,473", "0", "2.51", "3.70", "3.0", "7.0", "13.0", "40.25", "PASS"],
        ["GET", "/api/contracts/[id]/features", "2,044", "29", "2.07", "4.39", "4.0", "7.0", "9.0", "29.76", "404 Race"],
        ["GET", "/api/alerts/config", "1,979", "0", "2.01", "3.70", "3.0", "7.0", "8.0", "17.68", "PASS"],
        ["POST", "/api/auth/register [Expected 400]", "1,664", "0", "1.69", "4.07", "4.0", "7.0", "10.0", "28.41", "PASS"],
        ["POST", "/api/auth/login [Expected 401]", "1,589", "0", "1.61", "4.28", "4.0", "8.0", "10.0", "21.31", "PASS"],
        ["POST", "/api/contracts [Register]", "1,061", "0", "1.08", "7.50", "7.0", "11.0", "13.0", "23.05", "PASS"],
        ["DELETE", "/api/contracts/[id] [Cleanup]", "1,061", "0", "1.08", "3.87", "4.0", "5.0", "7.0", "22.75", "PASS"],
    ]

    t3 = doc.add_table(rows=len(t3_data)+1, cols=len(t3_headers))
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(t3_headers):
        t3.rows[0].cells[idx].text = name
    style_header_row(t3.rows[0])

    for row_idx, row_values in enumerate(t3_data):
        row = t3.rows[row_idx + 1]
        for col_idx, val in enumerate(row_values):
            row.cells[col_idx].text = val
        style_data_row(row, is_alternate=(row_idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 5. System Resource Utilization Matrix
    # -------------------------------------------------------------
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. System Resource Utilization Matrix")
    r_h5.font.color.rgb = RGBColor(26, 54, 93)

    t4_headers = ["Container / Process", "Component", "Initial RAM", "Peak RAM", "Final RAM", "RAM Delta", "Peak CPU"]
    t4_data = [
        ["prism_project-backend-1", "Spring Boot (JVM 21)", "705.3 MiB", "748.2 MiB", "746.8 MiB", "+41.5 MiB", "15% - 32%"],
        ["prism_project-postgres-1", "PostgreSQL 17 DB", "54.8 MiB", "55.4 MiB", "54.9 MiB", "+0.1 MiB", "3% - 8%"],
        ["prism_project-hardhat-node-1", "Local EVM RPC Node", "62.4 MiB", "62.4 MiB", "62.4 MiB", "0.0 MiB", "< 1%"],
        ["prism_project-static-analyzer-1", "FastAPI Analyzer", "48.2 MiB", "48.2 MiB", "48.2 MiB", "0.0 MiB", "< 1%"],
        ["Host System (Windows)", "OS, Docker & Locust", "12.8 GB", "13.2 GB", "13.0 GB", "+0.2 GB", "< 45% total"],
    ]

    t4 = doc.add_table(rows=len(t4_data)+1, cols=len(t4_headers))
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(t4_headers):
        t4.rows[0].cells[idx].text = name
    style_header_row(t4.rows[0])

    for row_idx, row_values in enumerate(t4_data):
        row = t4.rows[row_idx + 1]
        for col_idx, val in enumerate(row_values):
            row.cells[col_idx].text = val
        style_data_row(row, is_alternate=(row_idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 6. Service Level Objective (SLO) Evaluation Matrix
    # -------------------------------------------------------------
    h6 = doc.add_heading(level=1)
    r_h6 = h6.add_run("6. Service Level Objective (SLO) Evaluation Matrix")
    r_h6.font.color.rgb = RGBColor(26, 54, 93)

    t5_headers = ["SLO Metric", "Target Threshold", "Observed Baseline", "Observed 50k Run", "Margin / Safety Factor", "Compliance"]
    t5_data = [
        ["Error Rate (%)", "< 1.00%", "0.00%", "0.256%", "0.744% below limit", "PASS"],
        ["5xx Server Errors", "0.00%", "0.00%", "0.00%", "Zero crashes / leaks", "PASS"],
        ["Median Latency (p50)", "< 200 ms", "7.0 ms", "4.0 ms", "50x faster than SLO", "PASS"],
        ["95th Percentile (p95)", "< 500 ms", "14.0 ms", "9.0 ms", "55x faster than SLO", "PASS"],
        ["99th Percentile (p99)", "< 1500 ms", "18.0 ms", "11.0 ms", "136x faster than SLO", "PASS"],
        ["Emergency Tripwire", "Not triggered", "Safe", "Safe", "No tripwires breached", "PASS"],
    ]

    t5 = doc.add_table(rows=len(t5_data)+1, cols=len(t5_headers))
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, name in enumerate(t5_headers):
        t5.rows[0].cells[idx].text = name
    style_header_row(t5.rows[0])

    for row_idx, row_values in enumerate(t5_data):
        row = t5.rows[row_idx + 1]
        for col_idx, val in enumerate(row_values):
            row.cells[col_idx].text = val
        style_data_row(row, is_alternate=(row_idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 7. Concurrency Race & Root-Cause Explanation
    # -------------------------------------------------------------
    h7 = doc.add_heading(level=1)
    r_h7 = h7.add_run("7. Concurrency Race Condition Analysis")
    r_h7.font.color.rgb = RGBColor(26, 54, 93)

    p_race = doc.add_paragraph(
        "During the 50,000-request benchmark, 128 requests (0.256%) returned HTTP 404 across sub-entity endpoints "
        "(/api/contracts/{id}, /risk, /features, /telemetry). This occurred because the test suite exercised realistic "
        "CRUD lifecycles: 1,061 ephemeral test contracts were registered and subsequently cleaned up via DELETE. "
        "Due to concurrent greenlet execution, other workers queried contract details milliseconds after an ID was "
        "deleted from PostgreSQL. This is standard RESTful behavior for deleted resources, confirming that zero server-side "
        "exceptions (5xx) occurred."
    )
    p_race.style.font.size = Pt(10)

    # Save Document
    doc.save(output_path)
    print(f"Successfully generated DOCX report at: {output_path}")

if __name__ == "__main__":
    out_dir = os.path.join(os.path.dirname(__file__), "..", "reports")
    docx_file = os.path.join(out_dir, "PERFORMANCE_ANALYSIS_PARAMETERS_AND_MATRIX.docx")
    build_docx_report(os.path.abspath(docx_file))
